"""Convert defense_speech_5min.md to a polished .docx for the team.

Stripped-down markdown -> docx converter:
  - lines starting with "# "  -> Heading 1
  - lines starting with "## " -> Heading 2
  - lines starting with "> "  -> blockquote (italic)
  - separator "---"           -> skipped
  - everything else           -> body paragraphs

Bold (**...**) and inline code (`...`) inside body paragraphs are
rendered as runs with the appropriate formatting. The mix is enough
for the speech document — no tables, lists, or images required.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


SRC = Path(__file__).resolve().parent.parent / "docs" / "defense_speech_5min.md"
DST = Path(__file__).resolve().parent.parent / "docs" / "doklad_hr_ai_agent_5_min.docx"


_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_runs(paragraph, text: str) -> None:
    """Add inline runs to a paragraph, parsing **bold** and `code` markers."""
    parts = _INLINE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(11)
        else:
            paragraph.add_run(part)


def build(src: Path, dst: Path) -> None:
    md = src.read_text(encoding="utf-8")
    doc = Document()

    # Normal style: Arial 12 for safe rendering across machines.
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)

    for raw in md.splitlines():
        line = raw.rstrip()
        if not line:
            doc.add_paragraph()
            continue
        if line.strip() == "---":
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            run = p.add_run(line[2:].strip())
            run.italic = True
            continue
        # Body paragraph with inline formatting.
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_runs(p, line)

    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dst)
    print(f"Wrote: {dst}")


if __name__ == "__main__":
    build(SRC, DST)
